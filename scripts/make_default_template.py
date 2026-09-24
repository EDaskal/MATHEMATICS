"""Δημιουργεί το ενσωματωμένο πρότυπο (app/resources/default_template.docx).

Χρησιμοποιείται μόνο όταν η καθηγήτρια δεν έχει επιλέξει δικό της διαγώνισμα ως πρότυπο.
Βάση: το reference.docx του pandoc, με A4, περιθώρια 2 cm, «Σελίδα X από Y» στην κεφαλίδα.

    python scripts/make_default_template.py
"""

import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parents[1] / "app" / "resources" / "default_template.docx"
FONT = "Calibri"


def _field(paragraph, instr: str, bold: bool = True) -> None:
    run = paragraph.add_run()
    run.bold = bold
    for kind, text in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text == instr:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
            run._r.append(el)
        else:
            el = OxmlElement("w:t")
            el.text = text
            run._r.append(el)


def _font(style, size_pt: float | None = None) -> None:
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    if size_pt:
        style.font.size = Pt(size_pt)


def main() -> None:
    with tempfile.TemporaryDirectory() as td:
        ref = Path(td) / "ref.docx"
        subprocess.run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"], check=True)
        doc = Document(str(ref))

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
    sec.header_distance = sec.footer_distance = Cm(1)

    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        _font(st, 11)
        pf = st.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15

    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.add_run("Σελίδα ")
    _field(hp, "PAGE")
    hp.add_run(" από ")
    _field(hp, "NUMPAGES")

    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("")
    run.font.size = Pt(9)

    # το reference.docx του pandoc έχει δείγμα κειμένου· το σώμα αγνοείται από το pandoc, αλλά το καθαρίζουμε
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("OK", OUT)


if __name__ == "__main__":
    sys.exit(main())
