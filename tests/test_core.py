"""Δοκιμές πυρήνα (χωρίς κλήσεις στον Claude). Τρέχουν και στο GitHub πριν από το build."""

import json
import re
import zipfile
from pathlib import Path

import pytest

from app import points
from app.docx_build import BuildContext, build_docx, exam_markdown
from app.figures import FigureError, parse_spec, render_to_files, run_checks
from app.transcribe import EXERCISE_SCHEMA, normalize

FIX = Path(__file__).parent / "fixtures"


def load(name):
    return normalize(json.loads((FIX / name).read_text(encoding="utf-8")))


# ------------------------------------------------------------------ μονάδες

def test_points_presets_sum_to_100():
    for n in range(1, 7):
        assert sum(points.theme_points(n)) == 100


def test_split_even_remainder_first():
    assert points.split_even(25, 3) == [9, 8, 8]
    assert points.item_points(30, 2) == [15, 15]


def test_points_line():
    assert points.points_line([9, 8, 8]) == "Μονάδες 9 + 8 + 8"


# ------------------------------------------------------------------ σχήμα JSON

def test_schema_has_no_open_objects():
    """Κάθε αντικείμενο του σχήματος κλείνει με additionalProperties=false (απαίτηση structured outputs)."""

    def walk(node):
        if isinstance(node, dict):
            if node.get("type") == "object":
                assert node.get("additionalProperties") is False
                assert set(node["required"]) == set(node["properties"])
            for v in node.values():
                walk(v)
        elif isinstance(node, list):
            for v in node:
                walk(v)

    walk(EXERCISE_SCHEMA)


def test_normalize_limits_depth_and_choices():
    raw = {"stem": "x", "items": [{"label": "α)", "text": "t", "choices": {"columns": 3, "options": ["a", "b"]},
                                   "items": [{"label": "i", "text": "u", "choices": None,
                                              "items": [{"label": "1", "text": "v", "choices": None,
                                                         "items": [{"label": "deep", "text": "w", "choices": None, "items": []}]}]}]}],
           "figure": None}
    n = normalize(raw)
    assert n["items"][0]["label"] == "α"
    assert n["items"][0]["choices"]["columns"] == 2
    assert n["items"][0]["items"][0]["items"][0]["items"] == []


# ------------------------------------------------------------------ σχήματα

def test_figure_checks_pass_for_correct_intersection():
    ex = load("lines_figure.json")
    spec = parse_spec(ex["figure"]["spec_json"])
    assert run_checks(spec) == []


def test_figure_checks_catch_wrong_point():
    spec = {"axes": {"x": [-3, 3], "y": [-2, 7]}, "elements": [
        {"type": "line", "name": "ε", "through": [[-2, 0], [0, 2]]},
        {"type": "line", "name": "ζ", "through": [[0, 6], [2, 0]]},
        {"type": "point", "name": "M", "at": [2, 2]}],
        "checks": [{"type": "intersection", "lines": ["ε", "ζ"], "point": "M"}]}
    warnings = run_checks(spec)
    assert warnings and "(1.00, 3.00)" in warnings[0]


def test_render_figure_files(tmp_path):
    ex = load("lines_figure.json")
    info = render_to_files(ex["figure"]["spec_json"], tmp_path / "fig")
    svg = Path(info["svg"]).read_text(encoding="utf-8")
    assert "<svg" in svg and Path(info["png"]).stat().st_size > 1000
    assert "ε" in svg  # το κείμενο μένει κείμενο (επεξεργάσιμο)


# ------------------------------------------------------------------ Word

def _exam(tmp_path):
    lines = load("lines_figure.json")
    lines["figure_render"] = render_to_files(lines["figure"]["spec_json"], tmp_path / "f1")
    return {
        "title": "ΔΙΑΓΩΝΙΣΜΑ", "subtitle": "ΔΟΚΙΜΗ", "date": "",
        "themes": [
            {"exercises": [load("three_levels.json")], "points": [13, 12]},
            {"exercises": [load("crossref.json"), load("choices.json")], "points": [15, 15]},
            {"exercises": [lines], "points": [10, 10, 10]},
        ],
    }


def test_markdown_numbering_and_refs(tmp_path):
    ctx = BuildContext()
    md = exam_markdown(_exam(tmp_path), ctx)
    assert "**ΘΕΜΑ Α**" in md and "**Α1.**" in md and "**Α2.**" in md
    assert "**Β1.**" in md and "**Β2.**" in md and "**Γ3.**" in md
    assert "[[" not in md  # οι παραπομπές αντικαταστάθηκαν
    assert "του ερωτήματος (i)" in md  # Β1: το (α) έγινε i)
    assert "Μονάδες 13 + 12" in md
    assert re.search(r"^i\)\s", md, re.M)


def test_build_docx_end_to_end(tmp_path):
    from app.docx_build import PageSetup

    out = tmp_path / "exam.docx"
    res = build_docx(_exam(tmp_path), out, setup=PageSetup("Cambria", 12, "Β’ ΛΥΚΕΙΟΥ", "Δοκιμή Δοκιμίδου"))
    expected, found = res["math"]
    assert expected > 40 and expected == found, res["warnings"]
    with zipfile.ZipFile(out) as z:
        doc = z.read("word/document.xml").decode("utf-8")
        assert "svgBlip" in doc
        assert 'w:val="none"' in doc  # πίνακες χωρίς περιγράμματα
        assert "<m:oMath" in doc
        footer = z.read("word/footer_exam.xml").decode()
        assert "Β’ ΛΥΚΕΙΟΥ" in footer and "Επιμέλεια: Δοκιμή Δοκιμίδου" in footer and "<w:tab/>" in footer
        assert 'w:val="right"' in footer  # στάση στηλοθέτη στο δεξί περιθώριο
        header = z.read("word/header_exam.xml").decode()
        assert 'w:val="center"' in header and "PAGE" in header and "NUMPAGES" in header
        assert "rIdExamHdr" in doc and "rIdExamFtr" in doc
        styles = z.read("word/styles.xml").decode()
        assert 'w:ascii="Cambria"' in styles and 'w:sz w:val="24"' in styles
        assert "asciiTheme" not in styles
        ct = z.read("[Content_Types].xml").decode()
        assert "/word/header_exam.xml" in ct


def test_build_with_missing_template_fails(tmp_path):
    from app.docx_build import BuildError

    with pytest.raises(BuildError):
        build_docx({"themes": []}, tmp_path / "x.docx", template=str(tmp_path / "nope.docx"))


def test_header_footer_replace_template_with_first_page_header(tmp_path):
    """Πρότυπο με διαφορετική κεφαλίδα πρώτης σελίδας: αντικαθίσταται παντού από τη δική μας."""
    from docx import Document

    from app.docx_build import PageSetup, default_template

    tpl = tmp_path / "tpl.docx"
    d = Document(str(default_template()))
    sec = d.sections[0]
    sec.different_first_page_header_footer = True
    sec.first_page_header.paragraphs[0].text = "ΠΡΩΤΗ ΣΕΛΙΔΑ"
    d.save(str(tpl))
    out = tmp_path / "o.docx"
    build_docx({"title": "Τ", "themes": [{"exercises": [load("crossref.json")], "points": [50, 50]}]},
               out, template=str(tpl), setup=PageSetup("Georgia", 11, "Α", "Β"))
    with zipfile.ZipFile(out) as z:
        doc = z.read("word/document.xml").decode()
        sect = re.findall(r"<w:sectPr.*?</w:sectPr>", doc, re.S)[-1]
        assert "titlePg" not in sect
        assert sect.count("headerReference") == 1 and "rIdExamHdr" in sect
        assert 'w:sz w:val="22"' in z.read("word/header_exam.xml").decode()


def test_multi_page_prompt_and_crop_page():
    from app.transcribe import user_text

    assert "ΔΙΑΔΟΧΙΚΕΣ ΣΕΛΙΔΕΣ" in user_text(2) and "2 εικόνες" in user_text(2)
    assert "ΔΙΑΔΟΧΙΚΕΣ" not in user_text(1)
    raw = {"stem": "", "items": [], "figure": {"description": "d", "spec_json": "{}", "crop": [0, 0, 1, 1],
                                               "crop_page": 5, "confidence": "high"}}
    assert normalize(raw, n_pages=2)["figure"]["crop_page"] == 1


def test_versions_and_changelog():
    from app import __version__
    from app.updates import is_newer, parse_version

    assert parse_version("v1.10.2") == (1, 10, 2)
    assert is_newer("v1.2.0", "1.1.9") and not is_newer("1.1.0", "1.1.0") and is_newer("2.0", "1.9.9")
    import subprocess
    import sys

    out = subprocess.run([sys.executable, "scripts/changelog_section.py", __version__],
                         capture_output=True, text=True, encoding="utf-8")
    assert out.returncode == 0 and out.stdout.strip(), "Λείπει η τρέχουσα έκδοση από το CHANGELOG.md"


def test_settings_lists_and_types(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    from app.config import Settings

    s = Settings()
    s.update({"font_size": "11.5", "classes": ["Α", " ", "Β"], "check_updates": 0, "max_parallel": "4"})
    assert s.font_size == 11.5 and s.classes == ["Α", "Β"] and s.check_updates is False and s.max_parallel == 4
    s.remember("editors", "Χ")
    s.remember("editors", "Χ")
    assert s.editors == ["Χ"]


def test_update_check_parses_release(monkeypatch):
    from app import updates

    fake = {"tag_name": "v9.0.0", "body": "### Νέα\n- κάτι", "html_url": "https://github.com/x/y/releases/tag/v9.0.0",
            "published_at": "2030-01-02T10:00:00Z",
            "assets": [{"name": "notes.txt", "browser_download_url": "https://x/notes.txt", "size": 1},
                       {"name": "Diagonismata-Setup-9.0.0.exe", "browser_download_url": "https://x/Setup.exe", "size": 123}]}
    monkeypatch.setattr(updates, "_get_json", lambda url, timeout=10: fake)
    r = updates.check("EDaskal/MATHEMATICS", "1.1.0")
    assert r["newer"] and r["latest"] == "9.0.0" and r["url"] == "https://x/Setup.exe" and r["published"] == "2030-01-02"
    fake["tag_name"] = "v1.1.0"
    assert not updates.check("EDaskal/MATHEMATICS", "1.1.0")["newer"]


def test_download_rejects_insecure_url():
    import pytest as _pytest

    from app.updates import UpdateError, download

    with _pytest.raises(UpdateError):
        download("http://example.com/x.exe")


def test_math_font_in_settings_and_runs(tmp_path):
    from app.docx_build import PageSetup

    out = tmp_path / "m.docx"
    build_docx({"title": "Τ", "themes": [{"exercises": [load("three_levels.json")], "points": [50, 50]}]},
               out, setup=PageSetup("Cambria", 12, "Α", "Β", "STIX Two Math", 11))
    with zipfile.ZipFile(out) as z:
        settings = z.read("word/settings.xml").decode()
        doc = z.read("word/document.xml").decode()
    assert 'm:mathFont m:val="STIX Two Math"' in settings
    # κάθε τμήμα εξίσωσης έχει τη γραμματοσειρά και το μέγεθος μαθηματικών (11pt = 22 μισά σημεία)
    assert doc.count("<m:r>") > 0
    assert doc.count("<m:r>") == doc.count('<w:rFonts w:ascii="STIX Two Math"') - doc.count("<m:ctrlPr>")
    assert re.search(r'<m:r><w:rPr><w:rFonts w:ascii="STIX Two Math"[^>]*/><w:sz w:val="22"/>', doc)
    # σειρά μέσα στο m:r: m:rPr πριν από w:rPr
    assert not re.search(r"<m:r><w:rPr>(?:(?!</w:rPr>).)*</w:rPr><m:rPr>", doc)


def test_figure_font_from_settings_size_from_figure(tmp_path):
    """1.7.5: η γραμματοσειρά έρχεται από τις ρυθμίσεις, το μέγεθος όχι (όπως στο πρωτότυπο)."""
    ex = load("lines_figure.json")
    info = render_to_files(ex["figure"]["spec_json"], tmp_path / "f", font="STIX Two Math", size=20)
    svg = Path(info["svg"]).read_text(encoding="utf-8")
    assert "STIX Two Math" in svg and "20px" not in svg
    assert info["label_mode"] == "auto" and 8 <= info["label_pt"] <= 12


def test_figure_label_size_follows_original():
    from app.figures import label_size_pt

    # κεφαλαίο γράμμα = 5% του ύψους ενός σχήματος 6 cm → ~0,3 cm → ~12,5 pt
    assert 12 <= label_size_pt({"label_size": 0.05}, 7, 6) <= 13
    # μικρότερα γράμματα στο πρωτότυπο → μικρότερα κι εδώ· όρια 7–16
    assert label_size_pt({"label_size": 0.03}, 7, 6) < label_size_pt({"label_size": 0.05}, 7, 6)
    assert label_size_pt({"label_size": 0.4}, 16, 16) == 16
    assert label_size_pt({"label_size": 0.001}, 7, 6) == 7
    # ρητό μέγεθος από τα κουμπιά Α− / Α+ υπερισχύει
    assert label_size_pt({"label_size": 0.05, "label_pt": 9}, 7, 6) == 9
    assert label_size_pt({"label_size": "abc"}, 7, 7) == 10.5


def test_vector_labels_and_named_points(tmp_path):
    import json as _json

    from app.figures import _vec_parts

    assert _vec_parts(r"\vec{α}+\vec{β}") == ("α+β", [(0, 1), (2, 3)])
    assert _vec_parts(r"2\overrightarrow{AB}") == ("2AB", [(1, 3)])
    assert _vec_parts("\\\\vec{u}") == ("u", [(0, 1)])  # διπλό escape από JSON μέσα σε JSON
    assert _vec_parts("ε_1") == ("ε₁", [])
    spec = {"width_cm": 8, "label_size": 0.06, "axes": None, "elements": [
        {"type": "point", "name": "Α", "at": [0, 0], "label": "Α"},
        {"type": "point", "name": "Β", "at": [5, 0], "label": "Β"},
        {"type": "point", "name": "Γ", "at": [7, 3], "label": "Γ"},
        {"type": "point", "name": "Δ", "at": [2, 3], "label": "Δ", "show": False},
        {"type": "vector", "from": "Α", "to": "Β", "label": "α", "label_side": "right"},
        {"type": "vector", "from": "Α", "to": "Δ", "label": "AΔ", "marks": 2},
        {"type": "vector", "from": "Α", "to": "Γ", "label": r"\vec{α}+\vec{β}", "style": "dashed"},
        {"type": "vector", "from": "Β", "to": "Δ", "head": "mid"},
        {"type": "segment", "from": "Β", "to": "Γ", "marks": 2},
        {"type": "angle", "vertex": "Α", "from": "Β", "to": "Δ", "right": True},
        {"type": "line", "name": "ε", "through": ["Α", "Γ"]}],
        "checks": [{"type": "on_line", "point": "Γ", "line": "ε"}]}
    info = render_to_files(_json.dumps(spec, ensure_ascii=False), tmp_path / "v")
    assert info["warnings"] == [] and info["label_mode"] == "original"
    svg = Path(info["svg"]).read_text(encoding="utf-8")
    assert "α+β" in svg and "AΔ" in svg  # οι ετικέτες μένουν κείμενο (επεξεργάσιμο στο Word)
    with pytest.raises(FigureError):
        render_to_files(_json.dumps({"elements": [{"type": "vector", "from": "Χ", "to": [1, 1]}]}), tmp_path / "x")


def test_font_scan_detects_math_and_greek(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    from app.fonts import scan

    candidates = [Path("/usr/share/texmf/fonts/opentype/public/lm-math"), Path("/usr/share/fonts/truetype/dejavu")]
    dirs = [d for d in candidates if d.exists()]
    if not dirs:
        pytest.skip("δεν υπάρχουν δοκιμαστικές γραμματοσειρές σε αυτό το σύστημα")
    r = scan(refresh=True, dirs=dirs)
    assert r["count"] > 0 and r["text"]
    assert any(f["greek"] for f in r["text"])
    assert r["math"], "αναμενόταν τουλάχιστον μία γραμματοσειρά μαθηματικών"
    assert scan(dirs=dirs) == r  # από την cache


# ------------------------------------------------------------------ Corpus: εισαγωγή σε υπάρχον Word

CORPUS_EX = {
    "source_label": "", "stem": "Δίνονται τα σημεία $\\mathrm{Α}(1,0)$ και $\\mathrm{Β}(-2,3)$. Να βρείτε:",
    "items": [
        {"label": "i", "text": "την εξίσωση της ευθείας $\\mathrm{ΑΒ}$", "choices": None, "items": []},
        {"label": "ii", "text": "τον συντελεστή διεύθυνσης\n$$\\lambda=\\frac{y_2-y_1}{x_2-x_1}$$", "choices": None, "items": []},
        {"label": "iii", "text": "1. αν η ευθεία περνά από το $\\mathrm{Ο}$ (όπως στο [[i]])", "choices": None, "items": []},
    ],
    "closing": "", "figure": None, "uncertain": [],
}


def _corpus_doc(tmp_path):
    from app.corpus_doc import sample_document

    return sample_document(tmp_path / "corpus.docx", tmp_path)


def test_corpus_outline(tmp_path):
    from app.corpus_doc import outline_of

    hs = outline_of(_corpus_doc(tmp_path))
    assert [h["title"] for h in hs] == ["ΚΕΦΑΛΑΙΟ 1: ΕΥΘΕΙΑ", "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "ΓΕΝΙΚΗ ΜΟΡΦΗ", "ΚΕΦΑΛΑΙΟ 2: ΚΥΚΛΟΣ"]
    assert [h["depth"] for h in hs] == [0, 1, 1, 0]
    assert all(h["exercises"] == 0 for h in hs)


def test_corpus_insert_numbering_continues(tmp_path):
    from app.corpus_doc import CorpusDoc, insert_into, outline_of

    doc = _corpus_doc(tmp_path)
    key = {"title": "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "nth": 0}
    info = insert_into(doc, key, [CORPUS_EX, CORPUS_EX], tmp_path, make_backup=True)
    assert (info["first_label"], info["last_label"]) == ("1.", "2.")
    assert Path(info["backup"]).exists()
    info = insert_into(doc, key, [CORPUS_EX], tmp_path, make_backup=False)
    assert info["last_label"] == "3."
    hs = {h["title"]: h for h in outline_of(doc)}
    assert hs["ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ"]["exercises"] == 3
    assert hs["ΓΕΝΙΚΗ ΜΟΡΦΗ"]["exercises"] == 0  # οι ασκήσεις δεν έγιναν τίτλοι, η επόμενη ενότητα έμεινε ίδια
    with zipfile.ZipFile(doc) as z:
        xml = z.read("word/document.xml").decode()
    # 3 ασκήσεις × (2 εξισώσεις εκφώνησης + 1 + 1 ερωτημάτων + 1 σε δική της γραμμή + 1)
    assert xml.count("<m:oMath>") + xml.count("<m:oMathPara>") >= 3 * 6
    assert "[[" not in xml and "(a)" in xml  # παραπομπή [[i]] → (a)
    assert "1. αν η ευθεία" in xml  # το «1.» στην αρχή κειμένου δεν έγινε λίστα
    d = CorpusDoc(doc)
    texts = [d.text(p) for p in d.body if etree_local(p) == "p"]
    assert texts.index("ΓΕΝΙΚΗ ΜΟΡΦΗ") > max(i for i, t in enumerate(texts) if t.startswith("Δίνονται τα σημεία"))


def etree_local(el):
    from lxml import etree

    return etree.QName(el).localname


def test_corpus_unknown_section(tmp_path):
    from app.corpus_doc import CorpusError, insert_into

    with pytest.raises(CorpusError):
        insert_into(_corpus_doc(tmp_path), {"title": "ΔΕΝ ΥΠΑΡΧΕΙ", "nth": 0}, [CORPUS_EX], tmp_path)


def test_corpus_natural_order():
    from app.corpus_doc import natural_key

    names = ["1.6.10.png", "1.6.2.png", "1.6.1.1.png", "1.6.1.png", "1.6.1.10.png", "1.6.1.2.png"]
    assert sorted(names, key=natural_key) == ["1.6.1.png", "1.6.1.1.png", "1.6.1.2.png", "1.6.1.10.png", "1.6.2.png", "1.6.10.png"]


def test_corpus_schema_has_no_open_objects():
    from app.transcribe import CORPUS_SCHEMA

    def walk(n):
        if isinstance(n, dict):
            if n.get("type") == "object":
                assert n.get("additionalProperties") is False
            for v in n.values():
                walk(v)
        elif isinstance(n, list):
            for v in n:
                walk(v)

    walk(CORPUS_SCHEMA)
    from app.transcribe import corpus_text

    assert "{pages}" not in corpus_text(1)
    assert "ΔΙΑΔΟΧΙΚΕΣ" in corpus_text(2) and "ΔΙΑΔΟΧΙΚΕΣ" not in corpus_text(1)


# ------------------------------------------------------------------ κινητό (QR)

def test_phone_server_roundtrip():
    import base64
    import io
    import urllib.request

    from PIL import Image

    from app.phone import PhoneServer, qr_svg

    got = []
    srv = PhoneServer(lambda data, name, join: (got.append((name, join, data[:22])) or {"ok": True}))
    port = srv.start(port=0)
    try:
        base = f"http://127.0.0.1:{port}/m/{srv.token}"
        page = urllib.request.urlopen(base).read().decode("utf-8")
        assert "Φωτογράφιση" in page
        srv.target = "Corpus · 6. ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ"
        st = json.loads(urllib.request.urlopen(base + "/status").read())
        assert st["target"] == srv.target and st["count"] == 0
        buf = io.BytesIO()
        Image.new("RGB", (40, 30), "white").save(buf, format="JPEG")
        body = json.dumps({"name": "a.jpg", "join": True,
                           "data": "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode()}).encode()
        req = urllib.request.Request(base + "/upload", data=body, method="POST", headers={"Content-Type": "application/json"})
        res = json.loads(urllib.request.urlopen(req).read())
        assert res["ok"] and res["count"] == 1
        assert got == [("a.jpg", True, "data:image/jpeg;base64")]
        with pytest.raises(urllib.error.HTTPError) as e:  # λάθος κλειδί → 404
            urllib.request.urlopen(f"http://127.0.0.1:{port}/m/wrongtoken/status")
        assert e.value.code == 404
    finally:
        srv.stop()
    assert qr_svg("http://192.168.1.10:8766/m/x").startswith("<svg")


# ------------------------------------------------------------------ 1.4.1: στήλες, σύμβολα, αρίθμηση

def test_column_separator_normalized():
    from app.transcribe import normalize

    d = normalize({"stem": "x", "items": [{"label": "1", "text": "$A(1,2)$,⇥$B(3,4)$", "choices": None},
                                          {"label": "2", "text": "$A(1,2)$\t$\\theta$", "choices": None}]})
    assert d["items"][0]["text"] == "$A(1,2)$, ⇥ $B(3,4)$"
    assert d["items"][1]["text"] == "$A(1,2)$ ⇥ $\\theta$"  # το \theta δεν πειράζεται


def test_corpus_columns_become_aligned_tabs(tmp_path):
    from app.corpus_doc import insert_into

    ex = {"stem": "Να βρεθεί ο συντελεστής διεύθυνσης όταν:", "closing": "", "figure": None, "items": [
        {"label": "1", "text": "$\\mathrm{Α}(3,4)$, ⇥ $\\mathrm{Β}(6,7)$", "choices": None, "items": []},
        {"label": "2", "text": "$\\mathrm{Α}(3\\kappa+1,2\\kappa+1)$, ⇥ $\\mathrm{Β}(-1,-3)$", "choices": None, "items": []}]}
    doc = _corpus_doc(tmp_path)
    insert_into(doc, {"title": "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "nth": 0}, [ex], tmp_path, make_backup=False)
    with zipfile.ZipFile(doc) as z:
        xml = z.read("word/document.xml").decode()
    stops = re.findall(r'<w:tab w:val="left" w:pos="(\d+)"/>', xml)
    assert len(stops) == 2 and stops[0] == stops[1]  # ίδιος στηλοθέτης και στα δύο ερωτήματα
    assert xml.count("<w:tab/>") == 2 and "§TAB§" not in xml and "⇥" not in xml
    assert "−" not in xml and "<m:t>-" in xml  # μείον όπως το γράφει το Word


def test_exam_columns_and_minus(tmp_path):
    ex = load("three_levels.json")
    ex["items"][0]["text"] = "$\\mathrm{Α}(-1,2)$, ⇥ $\\mathrm{Β}(3,-4)$"
    out = tmp_path / "e.docx"
    build_docx({"title": "Τ", "themes": [{"exercises": [ex], "points": [50, 50]}]}, out)
    with zipfile.ZipFile(out) as z:
        xml = z.read("word/document.xml").decode()
    assert "  " in xml and "⇥" not in xml and "−" not in xml


# ------------------------------------------------------------------ 1.4.2: η οθόνη χωρίς παλιά cache

def test_ui_server_no_store_and_token(tmp_path, monkeypatch):
    import urllib.error
    import urllib.request

    monkeypatch.setenv("APPDATA", str(tmp_path))
    from app import __version__
    from app.bridge import Api
    from app.main import _start_ui_server

    url = _start_ui_server(Api())
    assert f"v={__version__}" in url and "&k=" in url
    r = urllib.request.urlopen(url)
    assert "no-store" in r.headers.get("Cache-Control", "")
    assert 'data-mode="corpus"' in r.read().decode("utf-8")
    base = url.split("/index.html")[0]
    token = url.split("&k=")[1]
    with pytest.raises(urllib.error.HTTPError) as e:  # χωρίς κλειδί: απαγορεύεται
        urllib.request.urlopen(urllib.request.Request(base + "/api/get_state", data=b"[]", method="POST"))
    assert e.value.code == 403
    req = urllib.request.Request(base + "/api/get_state", data=b"[]", method="POST",
                                 headers={"X-Api-Token": token, "Content-Type": "application/json"})
    assert json.loads(urllib.request.urlopen(req).read())["result"]["version"] == __version__
    # 1.7.1: άμυνα σε ξένες ιστοσελίδες (CSRF / DNS rebinding) και σε ιδιωτικές μεθόδους
    for headers, path, code in (
        ({"X-Api-Token": token, "Content-Type": "text/plain"}, "/api/get_state", 415),
        ({"X-Api-Token": token, "Content-Type": "application/json", "Host": "evil.example:80"}, "/api/get_state", 404),
        ({"X-Api-Token": token, "Content-Type": "application/json"}, "/api/_open_pdf", 404),
        ({"X-Api-Token": token, "Content-Type": "application/json"}, "/api/__init__", 404),
    ):
        with pytest.raises(urllib.error.HTTPError) as e:
            urllib.request.urlopen(urllib.request.Request(base + path, data=b"[]", method="POST", headers=headers))
        assert e.value.code == code, (headers, path)


# ------------------------------------------------------------------ 1.4.3: σωστή κάρτα δικτύου για το QR

def test_phone_prefers_real_network_over_virtualbox(monkeypatch):
    """Σταθερός υπολογιστής με καλώδιο σε δίκτυο 20.20.1.x και εγκατεστημένο VirtualBox (192.168.56.1)."""
    import socket as _socket

    from app import phone

    monkeypatch.setattr(phone, "_default_route_ip", lambda: "20.20.1.45")
    monkeypatch.setattr(phone, "_windows_adapters",
                        lambda: {"192.168.56.1": "VirtualBox Host-Only Network", "20.20.1.45": "Ethernet",
                                 "172.22.80.1": "vEthernet (WSL)", "127.0.0.1": "Loopback Pseudo-Interface 1"})
    monkeypatch.setattr(_socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", ("192.168.56.1", 0))])
    addrs = phone.lan_addresses()
    assert [ip for ip, _ in addrs] == ["20.20.1.45", "192.168.56.1", "172.22.80.1"]
    assert "Ethernet" in addrs[0][1] and "κύρια" in addrs[0][1]
    assert "όχι αυτή" in addrs[1][1]


def test_phone_virtualbox_last_without_adapter_names(monkeypatch):
    import socket as _socket

    from app import phone

    monkeypatch.setattr(phone, "_default_route_ip", lambda: None)
    monkeypatch.setattr(phone, "_windows_adapters", lambda: {})
    monkeypatch.setattr(_socket, "getaddrinfo", lambda *a, **k: [(0, 0, 0, "", ("192.168.56.1", 0)),
                                                                   (0, 0, 0, "", ("20.20.1.45", 0))])
    assert phone.lan_ips() == ["20.20.1.45", "192.168.56.1"]


# --------------------------------------------------------------------------- 1.5.0: ίδιες ασκήσεις

def test_mathtype_mtef_to_latex():
    from app.mathtype import mtef_to_latex

    # χειροποίητη MTEF 5: κεφαλίδα, γραμμή με [διάνυσμα(ΓΔ)] = κλάσμα(1/2), ρίζα(x^2)
    def ch(tf, code):
        return bytes([2, 0, 128 + tf]) + code.to_bytes(2, "little")

    def line(*objs):
        return bytes([1, 0]) + b"".join(objs) + b"\x00"

    def tmpl(sel, var, *slots):
        return bytes([3, 0, sel, var, 0]) + b"".join(slots) + b"\x00"

    null_line = bytes([1, 1])
    eq = (b"\x05\x01\x00\x06\x08DSMT6\x00\x00"
          + line(tmpl(31, 2, line(ch(5, 0x0393), ch(5, 0x0394))),  # \overrightarrow{ΓΔ}
                 ch(6, 0x3D),
                 tmpl(11, 0, line(ch(8, 0x31)), line(ch(8, 0x32))),  # \frac{1}{2}
                 ch(6, 0x2212),
                 tmpl(10, 0, line(ch(3, 0x78), tmpl(28, 0, null_line, line(ch(8, 0x32)))), null_line))
          + b"\x00")
    tex = mtef_to_latex(eq)
    assert tex == r"\overrightarrow{\Gamma \Delta}=\frac{1}{2}-\sqrt{x^{2}}"
    assert mtef_to_latex(b"\x1c\x00" + b"\x00" * 26 + eq) == tex  # με την κεφαλίδα OLE των 28 bytes
    assert mtef_to_latex(b"garbage") is None


def test_omml_to_latex():
    from lxml import etree

    from app.duplicates import omml_latex

    M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
    xml = (f'<m:oMath {M}><m:acc><m:accPr><m:chr m:val="⃗"/></m:accPr><m:e><m:r><m:t>ΑΒ</m:t></m:r></m:e></m:acc>'
           '<m:r><m:t>=</m:t></m:r><m:f><m:num><m:r><m:t>1</m:t></m:r></m:num><m:den><m:r><m:t>2</m:t></m:r></m:den></m:f>'
           '<m:sSup><m:e><m:r><m:t>x</m:t></m:r></m:e><m:sup><m:r><m:t>2</m:t></m:r></m:sup></m:sSup>'
           '<m:r><m:rPr><m:scr m:val="double-struck"/></m:rPr><m:t>R</m:t></m:r></m:oMath>')
    tex = omml_latex(etree.fromstring(xml))
    assert tex.startswith(r"\overrightarrow{ΑΒ}=\frac{1}{2}{x}^{2}")
    assert r"\mathbb{R}" in tex


def test_canon_ignores_typography():
    from app.duplicates import canon, mask_numbers, pretty

    a = r"Δίνονται τα σημεία $\mathrm{Α}(1,0)$ και $\mathrm{Β}(-2,3)$. Να βρείτε το $\vec{\alpha}\cdot\vec{\beta}$, $x\in\mathbb{R}$"
    b = r"Δίνονται τα σημεία $Α\left(1, 0\right)$ και $B(−2,3)$ · Να βρείτε το $\overrightarrow{α}⋅\overrightarrow{β}$, $x \in ℝ$"
    assert canon(a) == canon(b)  # \mathrm, \left, κενά, −/-, λατινικό B, \vec/\overrightarrow, ℝ
    assert canon(r"$\frac{1}{2}$") == canon("1/2") and canon(r"$\sin x$") == canon(r"$\mathrm{ημ}x$")
    assert mask_numbers(canon("Α(1,0) και Β(-2,3)")) == mask_numbers(canon("Α(4,7) και Β(-5,9)"))
    assert pretty(r"$\overrightarrow{\mathrm{ΑΒ}}=(1,\frac{1}{2})$") == "ΑΒ⃗=(1,1/2)"


def test_duplicates_in_corpus_document(tmp_path):
    import copy

    from app.corpus_doc import insert_into
    from app.duplicates import check, index_of

    doc = _corpus_doc(tmp_path)
    other = {"source_label": "", "stem": "Να βρείτε την απόσταση του σημείου $\\mathrm{Μ}(3,4)$ από την ευθεία "
             "$\\varepsilon: 3x+4y-5=0$ και το εμβαδόν του τριγώνου που σχηματίζει η $\\varepsilon$ με τους άξονες.",
             "items": [], "closing": "", "figure": None, "uncertain": []}
    insert_into(doc, {"title": "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "nth": 0}, [CORPUS_EX, other], tmp_path, make_backup=False)
    idx = index_of(doc)
    labels = [e["label"] for e in idx["exercises"]]
    assert labels == ["1.", "2."]
    assert idx["exercises"][0]["parts"][1]["label"] == "a."  # τα ερωτήματα με την αρίθμηση του Word

    same = copy.deepcopy(CORPUS_EX)
    same["stem"] = same["stem"].replace("$\\mathrm{Α}(1,0)$", "$Α(1, 0)$")  # μόνο τυπογραφική διαφορά
    numbers = copy.deepcopy(other)
    numbers["stem"] = numbers["stem"].replace("(3,4)", "(6,2)").replace("3x+4y-5", "5x-12y+7")
    part = {"source_label": "", "stem": "Δίνεται η ευθεία $\\zeta: y=2x-1$. Να βρείτε:", "closing": "", "figure": None,
            "uncertain": [], "items": [{"label": "a", "text": "τα σημεία τομής της με τους άξονες", "choices": None, "items": []},
                                       copy.deepcopy(CORPUS_EX["items"][1])]}
    fresh = {"source_label": "", "stem": "Να λύσετε την εξίσωση $x^2-5x+6=0$ και να ελέγξετε τις λύσεις στο τετράδιο.",
             "items": [], "closing": "", "figure": None, "uncertain": []}
    new = [same, numbers, part, fresh, copy.deepcopy(fresh)]

    res = check(new, idx["exercises"], "normal")
    assert res[0][0]["level"] == "same" and res[0][0]["label"] == "1."
    assert res[1][0]["level"] == "numbers" and res[1][0]["label"] == "2." and res[1][0]["diffs"]
    assert res[2][0]["level"] == "part" and res[2][0]["part"] == "1" and res[2][0]["other_part"] == "b."
    assert res[3] == []
    assert res[4][0]["kind"] == "batch" and res[4][0]["index"] == 3 and res[4][0]["level"] == "same"
    # «μόνο ίδιες»: χωρίς «άλλα νούμερα» · ανενεργός: τίποτα
    low = check(new, idx["exercises"], "low")
    assert low[0][0]["level"] == "same" and low[1] == []
    assert check(new, idx["exercises"], "off") == [[]] * len(new)


def test_bridge_corpus_check(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    from app.bridge import Api
    from app.corpus_doc import insert_into

    doc = _corpus_doc(tmp_path)
    insert_into(doc, {"title": "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "nth": 0}, [CORPUS_EX], tmp_path, make_backup=False)
    api = Api()
    r = api.corpus_check(str(doc), [CORPUS_EX], "high")
    assert r["ok"] and r["existing"] == 1 and r["matches"][0][0]["level"] == "same"
    assert api._settings.dup_check == "high"
    assert api.corpus_check(str(tmp_path / "missing.docx"), [CORPUS_EX])["ok"] is False


def test_audit_whole_document_and_second_file(tmp_path, monkeypatch):
    import copy

    from app.corpus_doc import insert_into, sample_document
    from app.duplicates import audit, index_of

    a_dir, b_dir = tmp_path / "a", tmp_path / "b"
    a_dir.mkdir()
    b_dir.mkdir()
    doc_a = sample_document(a_dir / "a.docx", a_dir)
    doc_b = sample_document(b_dir / "b.docx", b_dir)
    other = {"source_label": "", "stem": "Να βρείτε την απόσταση του σημείου $\\mathrm{Μ}(3,4)$ από την ευθεία "
             "$\\varepsilon: 3x+4y-5=0$ και το εμβαδόν του τριγώνου που σχηματίζει η $\\varepsilon$ με τους άξονες.",
             "items": [], "closing": "", "figure": None, "uncertain": []}
    numbers = copy.deepcopy(other)
    numbers["stem"] = numbers["stem"].replace("(3,4)", "(6,2)")
    fresh = {"source_label": "", "stem": "Να λύσετε την εξίσωση $x^2-5x+6=0$ και να ελέγξετε τις λύσεις στο τετράδιο.",
             "items": [], "closing": "", "figure": None, "uncertain": []}
    key = {"title": "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "nth": 0}
    insert_into(doc_a, key, [CORPUS_EX, other, copy.deepcopy(CORPUS_EX), numbers], a_dir, make_backup=False)
    insert_into(doc_b, key, [fresh, copy.deepcopy(other)], b_dir, make_backup=False)
    ea, eb = index_of(doc_a)["exercises"], index_of(doc_b)["exercises"]

    r = audit(ea, None, "normal")
    got = {(p["level"], ea[p["a"]]["label"], ea[p["b"]]["label"]) for p in r["pairs"]}
    assert ("same", "1.", "3.") in got and ("numbers", "2.", "4.") in got
    assert all(not p["cross"] and p["a"] < p["b"] for p in r["pairs"])

    r = audit(ea, eb, "low")
    cross = [p for p in r["pairs"] if p["cross"]]
    assert [(eb[p["b"]]["label"], ea[p["a"]]["label"], p["level"]) for p in cross] == [("2.", "2.", "same")]
    assert {p["level"] for p in r["pairs"]} == {"same"}  # «μόνο ίδιες»

    monkeypatch.setenv("APPDATA", str(tmp_path / "appdata"))
    from app.bridge import Api

    res = Api().corpus_audit(str(doc_a), str(doc_b), "normal")
    assert res["ok"] and res["b"]["name"] == "b.docx" and len(res["a"]["exercises"]) == 4
    assert Api().corpus_audit(str(doc_a), str(doc_a))["ok"] is False  # ίδιο αρχείο


# --------------------------------------------------------------------------- 1.7.0: θέση και μετακίνηση ασκήσεων

def _ex(text):
    return {"source_label": "", "stem": text, "closing": "", "figure": None, "uncertain": [],
            "items": [{"label": "a", "text": f"ερώτημα της «{text[:12]}» με $x={len(text)}$", "choices": None, "items": []}]}


def _stems(doc):
    from app.duplicates import index_of

    return [(e["section"], e["label"], e["parts"][0]["md"][:6]) for e in index_of(doc)["exercises"]]


def test_move_and_place_exercises(tmp_path):
    import pytest

    from app.corpus_doc import CorpusError, file_sha, insert_into, move_in, undo_change
    from app.duplicates import index_of

    doc = _corpus_doc(tmp_path)
    a_key, b_key = {"title": "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "nth": 0}, {"title": "ΓΕΝΙΚΗ ΜΟΡΦΗ", "nth": 0}
    insert_into(doc, a_key, [_ex("ΑΑΑ πρώτη άσκηση της Α"), _ex("ΒΒΒ δεύτερη άσκηση της Α"),
                             _ex("ΓΓΓ τρίτη άσκηση της Α")], tmp_path, make_backup=False)
    insert_into(doc, b_key, [_ex("ΔΔΔ πρώτη της Β"), _ex("ΕΕΕ δεύτερη της Β")], tmp_path, make_backup=False)
    ex = {e["parts"][0]["md"][:3]: e for e in index_of(doc)["exercises"]}

    # μέσα στην ίδια ενότητα: η ΓΓΓ πριν από την ΑΑΑ
    r = move_in(doc, ex["ΓΓΓ"], ex["ΑΑΑ"], "before")
    assert (r["moved_from"], r["moved_to"]) == ("3.", "1.")
    assert [s[2][:3] for s in _stems(doc)][:3] == ["ΓΓΓ", "ΑΑΑ", "ΒΒΒ"]
    assert any("1.–2. → 2.–3." in t for t in r["renumbered"]) and Path(r["backup"]).exists()
    # αναίρεση: μόνο αν το αρχείο δεν άλλαξε, και επαναφέρει ακριβώς το αντίγραφο
    sha_after = file_sha(doc)
    undo_change(doc, r["backup"], r["sha"])
    assert file_sha(doc) != sha_after and [s[2][:3] for s in _stems(doc)][:3] == ["ΑΑΑ", "ΒΒΒ", "ΓΓΓ"]
    with pytest.raises(CorpusError):
        undo_change(doc, r["backup"], "λάθος")

    # σε άλλη ενότητα: η ΒΒΒ μετά την ΔΔΔ → παίρνει την αρίθμηση της ενότητας Β
    ex = {e["parts"][0]["md"][:3]: e for e in index_of(doc)["exercises"]}
    r = move_in(doc, ex["ΒΒΒ"], ex["ΔΔΔ"], "after", make_backup=False)
    st = _stems(doc)
    assert [(s[0], s[1], s[2][:3]) for s in st] == [
        ("ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "1.", "ΑΑΑ"), ("ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "2.", "ΓΓΓ"),
        ("ΓΕΝΙΚΗ ΜΟΡΦΗ", "1.", "ΔΔΔ"), ("ΓΕΝΙΚΗ ΜΟΡΦΗ", "2.", "ΒΒΒ"), ("ΓΕΝΙΚΗ ΜΟΡΦΗ", "3.", "ΕΕΕ")]
    assert r["section"] == "ΓΕΝΙΚΗ ΜΟΡΦΗ" and "Αρίθμηση: συνεχόμενη" in r["checks"][-1]
    idx = index_of(doc)["exercises"]
    assert idx[3]["parts"][1]["label"] == "a."  # και τα ερωτήματα στην αρίθμηση της νέας ενότητας

    # ήδη στη θέση του → μήνυμα, χωρίς αλλαγή
    ex = {e["parts"][0]["md"][:3]: e for e in idx}
    before = file_sha(doc)
    with pytest.raises(CorpusError, match="ήδη"):
        move_in(doc, ex["ΒΒΒ"], ex["ΔΔΔ"], "after")
    assert file_sha(doc) == before

    # νέα άσκηση σε συγκεκριμένη θέση + μία στο τέλος της ενότητας
    new1, new2 = _ex("ΖΖΖ νέα μετά την ΑΑΑ"), _ex("ΗΗΗ νέα στο τέλος")
    new1["place"] = {"side": "after", "fp": ex["ΑΑΑ"]["fp"], "label": ex["ΑΑΑ"]["label"]}
    info = insert_into(doc, a_key, [new1, new2], tmp_path, make_backup=False)
    assert info["labels"] == ["2.", "4."] and any("2. → 3." in t for t in info["renumbered"])
    assert [s[2][:3] for s in _stems(doc) if s[0] == "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ"] == ["ΑΑΑ", "ΖΖΖ", "ΓΓΓ", "ΗΗΗ"]
    # θέση που δεν υπάρχει πια → καθαρό μήνυμα, το αρχείο δεν αλλάζει
    bad = _ex("ΘΘΘ")
    bad["place"] = {"side": "after", "fp": "000000000000", "label": "9."}
    before = file_sha(doc)
    with pytest.raises(CorpusError, match="δεν βρέθηκε"):
        insert_into(doc, a_key, [bad], tmp_path, make_backup=False)
    assert file_sha(doc) == before


def test_numbering_issues_and_summary():
    from app.corpus_doc import numbering_issues

    exs = [{"label": f"4.{n}", "section": "4."} for n in (1, 2, 4)] + [{"label": "5.1", "section": "5."}]
    assert numbering_issues(exs) == ["4.: μετά την 4.2 ακολουθεί η 4.4"]


def test_exercise_without_stem_is_kept(tmp_path):
    from app.corpus_doc import insert_into, outline_of
    from app.duplicates import index_of

    doc = _corpus_doc(tmp_path)
    key = {"title": "ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ", "nth": 0}
    no_stem = dict(_ex("x"), stem="")
    info = insert_into(doc, key, [_ex("ΑΑΑ πρώτη"), no_stem, _ex("ΓΓΓ τρίτη")], tmp_path, make_backup=False)
    assert info["labels"] == ["1.", "2.", "3."]
    # η επόμενη εισαγωγή δεν τη σβήνει ως «κενή θέση»
    info = insert_into(doc, key, [_ex("ΔΔΔ τέταρτη")], tmp_path, make_backup=False)
    assert info["removed_placeholders"] == 0 and info["labels"] == ["4."]
    assert [e["label"] for e in index_of(doc)["exercises"]] == ["1.", "2.", "3.", "4."]
    assert {h["title"]: h["exercises"] for h in outline_of(doc)}["ΕΞΙΣΩΣΗ ΕΥΘΕΙΑΣ"] == 4


def _sample_pdf(pages: int = 3) -> bytes:
    import io

    import matplotlib

    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages

    buf = io.BytesIO()
    with PdfPages(buf) as pdf:
        for k in range(1, pages + 1):
            fig = plt.figure(figsize=(8.27, 11.69))
            fig.text(0.1, 0.9, f"{k}. Να λύσετε την εξίσωση $x^2-{k}=0$.", fontsize=16)
            pdf.savefig(fig)
            plt.close(fig)
    return buf.getvalue()


def test_pdf_render_pages_and_errors():
    from app.pdf_pages import PdfError, render

    doc = render(_sample_pdf(3))
    assert doc.total == 3 and len(doc.pages) == 3 and not doc.truncated
    assert all(max(im.size) == 2000 for im in doc.pages)
    assert doc.pages[0].convert("L").getextrema()[0] < 100  # υπάρχει κείμενο (σκούρα pixels)
    part = render(_sample_pdf(3), max_pages=2)
    assert part.total == 3 and len(part.pages) == 2 and part.truncated
    for bad, msg in ((b"", "κενό"), (b"GIF89a....", "δεν είναι PDF"), (b"%PDF-1.4\ngarbage", "δεν ανοίγει")):
        with pytest.raises(PdfError, match=msg):
            render(bad)


def test_phone_server_pdf_upload():
    import urllib.parse
    import urllib.request

    from app.phone import PhoneServer

    got = []
    srv = PhoneServer(lambda *a: {"ok": True}, lambda raw, name: (got.append((name, raw[:5])) or {"ok": True, "pages": 3}))
    port = srv.start(port=0)
    try:
        base = f"http://127.0.0.1:{port}/m/{srv.token}"

        def post(body: bytes, name: str):
            req = urllib.request.Request(base + "/upload-pdf", data=body, method="POST",
                                         headers={"Content-Type": "application/pdf",
                                                  "X-File-Name": urllib.parse.quote(name)})
            return json.loads(urllib.request.urlopen(req).read())

        res = post(_sample_pdf(1), "Ασκήσεις κεφ. 2.pdf")
        assert res["ok"] and res["pages"] == 3 and res["count"] == 1
        assert got == [("Ασκήσεις κεφ. 2.pdf", b"%PDF-")]
        with pytest.raises(urllib.error.HTTPError) as e:  # όχι PDF → 400, δεν φτάνει στην εφαρμογή
            post(b"\x89PNG\r\n\x1a\n" + b"0" * 100, "x.pdf")
        assert e.value.code == 400 and len(got) == 1
        with pytest.raises(urllib.error.HTTPError) as e:  # λάθος κλειδί → 404
            urllib.request.urlopen(urllib.request.Request(f"http://127.0.0.1:{port}/m/wrong/upload-pdf",
                                                          data=b"%PDF-1.4", method="POST"))
        assert e.value.code == 404
    finally:
        srv.stop()


def test_bridge_pdf_open_and_take(tmp_path, monkeypatch):
    import base64

    monkeypatch.setenv("APPDATA", str(tmp_path))
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    from app.bridge import Api

    api = Api()
    events = []
    api.emit = lambda name, data: events.append((name, data))
    res = api.add_pdf("data:application/pdf;base64," + base64.b64encode(_sample_pdf(3)).decode(), "Φυλλάδιο.pdf")
    assert res["ok"] and res["total"] == 3 and [p["n"] for p in res["pages"]] == [1, 2, 3]
    assert res["pages"][0]["thumb"].startswith("data:image/jpeg;base64,")
    took = api.pdf_take(res["pdf_id"], [3, 1, 99])
    assert took["ok"] and [p["n"] for p in took["pages"]] == [3, 1]
    assert took["pages"][0]["name"] == "Φυλλάδιο.pdf — σελ. 3"
    assert all(p["image_id"] in api._images for p in took["pages"])
    assert api.pdf_preview(res["pdf_id"], 2)["ok"] and not api.pdf_preview(res["pdf_id"], 9)["ok"]
    assert not api.pdf_take("άγνωστο", [1])["ok"]
    assert not api.add_pdf("data:application/pdf;base64," + base64.b64encode(b"not a pdf").decode(), "x.pdf")["ok"]
    # από το κινητό: γεγονός «phone_pdf» με τις μικρογραφίες
    ph = api._phone_pdf(_sample_pdf(2), "kinito.pdf")
    assert ph == {"ok": True, "pages": 2, "shown": 2}
    assert events and events[-1][0] == "phone_pdf" and events[-1][1]["name"] == "kinito.pdf"


def test_update_only_from_own_github_releases(tmp_path, monkeypatch):
    from app.updates import is_release_asset

    repo = "EDaskal/MATHEMATICS"
    ok = "https://github.com/EDaskal/MATHEMATICS/releases/download/v1.7.1/Diagonismata-Setup-1.7.1.exe"
    assert is_release_asset(ok, repo)
    for bad in ("http://github.com/EDaskal/MATHEMATICS/releases/download/v1/x.exe",
                "https://github.com.evil.com/EDaskal/MATHEMATICS/releases/download/v1/x.exe",
                "https://evil.com/EDaskal/MATHEMATICS/releases/download/v1/x.exe",
                "https://github.com/Other/MATHEMATICS/releases/download/v1/x.exe",
                "https://github.com/EDaskal/MATHEMATICS/releases/download/v1/x.zip",
                "https://user@github.com/EDaskal/MATHEMATICS/releases/download/v1/x.exe",
                "https://github.com/EDaskal/MATHEMATICS/releases/download/../../x.exe"):
        assert not is_release_asset(bad, repo), bad
    monkeypatch.setenv("APPDATA", str(tmp_path))
    from app.bridge import Api

    api = Api()
    started = []
    monkeypatch.setattr(api, "_update_worker", lambda url: started.append(url))
    assert not api.install_update("https://evil.com/x.exe")["ok"]  # χωρίς προηγούμενο έλεγχο: τίποτα
    api._update_url = ok
    assert api.install_update("https://evil.com/x.exe")["ok"]  # αγνοεί τη διεύθυνση της οθόνης
    import time

    time.sleep(0.2)
    assert started == [ok]


def test_open_file_only_documents(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    from app import bridge

    opened = []
    monkeypatch.setattr(bridge, "open_path", lambda p: opened.append(str(p)))
    api = bridge.Api()
    doc = tmp_path / "Διαγώνισμα.docx"
    doc.write_bytes(b"x")
    exe = tmp_path / "evil.exe"
    exe.write_bytes(b"MZ")
    assert api.open_file(str(doc))["ok"]
    assert not api.open_file(str(exe))["ok"]
    assert not api.open_file(str(tmp_path / "missing.docx"))["ok"]
    assert not api.open_file("https://evil.example/a.docx")["ok"]
    assert api.open_folder(str(doc))["ok"]  # ο φάκελος του αρχείου
    assert opened == [str(doc.resolve()), str(tmp_path.resolve())]


# ------------------------------------------------------------------ 1.7.2: χρήση συνδρομής Claude

RATE_RAW = {"status": "allowed_warning", "resetsAt": 1793286000, "rateLimitType": "seven_day", "overageStatus": "rejected",
            "overageDisabledReason": "org_level_disabled", "isUsingOverage": False,
            "unifiedWindows": {"five_hour": {"utilization": 0.25, "resetsAt": 1793217000},
                               "seven_day": {"utilization": 0.72, "resetsAt": 1793286000}}}


def test_usage_parse_and_persist(tmp_path):
    from app.usage import UsageTracker, parse_limits

    lim = parse_limits(RATE_RAW)
    assert lim["status"] == "allowed_warning" and lim["limit_type"] == "seven_day"
    assert lim["windows"]["five_hour"] == {"utilization": 0.25, "resets_at": 1793217000}
    assert parse_limits({"weird": 1})["windows"] == {} and parse_limits(None)["status"] == ""  # ανεκτικό
    t = UsageTracker(tmp_path / "usage.json")
    seen = []
    t.on_change = seen.append
    t.record_limits(RATE_RAW)
    t.record_result(images=2, usage={"input_tokens": 10, "output_tokens": 500, "cache_read_input_tokens": 6000},
                    cost_usd=0.02, seconds=4.2, model="claude-opus-5-5")
    snap = t.snapshot()
    assert snap["session"]["images"] == 2 and snap["today"]["output_tokens"] == 500 and snap["week"]["calls"] == 1
    assert snap["model"] == "claude-opus-5-5" and len(seen) == 2
    t2 = UsageTracker(tmp_path / "usage.json")  # νέο άνοιγμα της εφαρμογής: τα όρια και το «σήμερα» μένουν
    s2 = t2.snapshot()
    assert s2["limits"]["windows"]["seven_day"]["utilization"] == 0.72 and s2["today"]["images"] == 2
    assert s2["session"]["calls"] == 0
    assert "Ξαναδουλεύει" in t2.limit_message()


def test_usage_estimate_per_image(tmp_path):
    from app.usage import UsageTracker

    t = UsageTracker(tmp_path / "u.json")

    def lim(u, reset=1000):
        return {"status": "allowed", "unifiedWindows": {"five_hour": {"utilization": u, "resetsAt": reset}}}

    t.record_limits(lim(0.10))  # βάση
    for _ in range(4):
        t.record_result(images=1)
    t.record_limits(lim(0.11))
    assert t.per_image is None  # μόλις 1% αλλαγή: όχι ακόμα εκτίμηση
    t.record_limits(lim(0.14))
    assert abs(t.per_image - 0.01) < 1e-9 and t.per_image_samples == 4  # 4% σε 4 εικόνες
    t.record_limits(lim(0.02, reset=2000))  # νέο 5ωρο: νέα βάση, η εκτίμηση μένει
    assert abs(t.per_image - 0.01) < 1e-9


def test_bridge_usage_state_and_api_mode(tmp_path, monkeypatch):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    from app import usage
    from app.bridge import Api

    monkeypatch.setattr(usage, "tracker", usage.UsageTracker(tmp_path / "usage.json"))
    import app.bridge as br

    api = Api()
    usage.tracker.on_change = lambda snap: None
    st = api.usage_state()
    assert st["provider"] == "agent_sdk" and st["auto"] is True and st["limits"] is None
    api._settings.provider = "api"
    r = api.usage_refresh()
    assert not r["ok"] and "συνδρομή" in r["error"]
    assert br is not None


# ------------------------------------------------------------------ 1.7.6: χρώματα, βελάκια, έλεγχος σχημάτων

def _png(w=400, h=300):
    import io as _io

    from PIL import Image

    buf = _io.BytesIO()
    Image.new("RGB", (w, h), "white").save(buf, format="PNG")
    return buf.getvalue()


def test_figure_colors_and_bw(tmp_path):
    import json as _json

    spec = {"axes": None, "elements": [
        {"type": "vector", "from": [0, 0], "to": [3, 1], "label": "x", "color": "red"},
        {"type": "segment", "from": [0, 0], "to": [1, 3], "color": "#1565C0", "emphasis": True},
        {"type": "polygon", "points": [[0, 0], [3, 1], [1, 3]], "fill": "green"}]}
    col = Path(render_to_files(_json.dumps(spec), tmp_path / "c", colors="original")["svg"]).read_text(encoding="utf-8").lower()
    bw = Path(render_to_files(_json.dumps(spec), tmp_path / "b", colors="bw")["svg"]).read_text(encoding="utf-8").lower()
    assert "#c62828" in col and "#1565c0" in col
    assert "#c62828" not in bw and "#1565c0" not in bw and "#2e7d32" not in bw


def test_vec_arrow_clears_tall_letters():
    from matplotlib.font_manager import FontProperties

    from app.figures import _ink_top

    prop = FontProperties(family="serif", size=12)
    assert _ink_top("δ", prop) > _ink_top("α", prop) + 1  # το δ είναι ψηλότερο· το βελάκι μπαίνει πιο πάνω
    assert _ink_top("AB", prop) > _ink_top("x", prop)


def test_needs_check_and_grid():
    import json as _json

    from PIL import Image

    from app.figure_check import exercise_text, grid_crop, needs_check

    vec = {"spec_json": _json.dumps({"axes": {"x": [0, 5], "y": [0, 5]}, "elements": [{"type": "vector", "from": [0, 0], "to": [1, 1]}]}), "confidence": "high"}
    graph = {"spec_json": _json.dumps({"axes": {"x": [0, 5], "y": [0, 5]}, "elements": [{"type": "function", "expr": "x"}]}), "confidence": "high"}
    geo = {"spec_json": _json.dumps({"axes": None, "elements": []}), "confidence": "high"}
    assert needs_check(vec) and needs_check(geo) and not needs_check(graph)
    assert needs_check({**graph, "confidence": "low"}) and not needs_check(None)
    g = Image.open(__import__("io").BytesIO(grid_crop(_png(), [0.2, 0.2, 0.6, 0.5])))
    assert max(g.size) > 1400  # μεγεθυμένη περικοπή + περιθώριο για τους αριθμούς
    assert "α) Να βρείτε" in exercise_text({"stem": "Δίνεται", "items": [{"label": "α", "text": "Να βρείτε", "items": []}]})


class _FakeProvider:
    def __init__(self, answers):
        self.answers = list(answers)
        self.calls = []

    async def extract_json(self, system, user_text, images, schema, count_images=None):
        self.calls.append({"text": user_text, "n": len(images), "count": count_images})
        return self.answers.pop(0)


def test_check_figure_replaces_or_keeps():
    import asyncio
    import json as _json

    from app.figure_check import check_figure

    first = {"spec_json": _json.dumps({"axes": None, "elements": [{"type": "vector", "from": [0, 0], "to": [1, 1]}]}),
             "crop": [0, 0, 1, 1], "crop_page": 0, "confidence": "low", "description": "d"}
    good = _json.dumps({"axes": None, "elements": [{"type": "vector", "from": [1, 1], "to": [0, 0], "label": "α"}]})
    fp = _FakeProvider([{"notes": "α: από (1,1) προς (0,0)", "changes": "φορά α", "description": "νέο", "confidence": "high", "spec_json": good}])
    out = asyncio.run(check_figure(fp, [_png()], first, "Να εκφράσετε το α", "κανόνες"))
    assert out["spec_json"] == good and out["checked"] and out["confidence"] == "high"
    assert fp.calls[0]["n"] == 2 and fp.calls[0]["count"] == 0  # πλέγμα + πρώτη σχεδίαση, δεν μετρά ως ανάγνωση
    assert "Να εκφράσετε το α" in fp.calls[0]["text"]
    bad = _FakeProvider([{"notes": "", "changes": "", "description": "", "confidence": "high", "spec_json": "{όχι json"}])
    assert asyncio.run(check_figure(bad, [_png()], first, "", "κ")) is first


def test_read_figure_from_solution():
    import asyncio
    import json as _json

    from app.figure_check import read_figure

    spec = _json.dumps({"axes": None, "elements": [{"type": "segment", "from": [0, 0], "to": [2, 1]}]})
    fp = _FakeProvider([
        {"found": True, "crop_page": 1, "crop": [0.5, 0.1, 0.9, 0.4], "description": "τρίγωνο της λύσης", "why": ""},
        {"notes": "", "changes": "", "description": "τρίγωνο", "confidence": "high", "spec_json": spec}])
    fig = asyncio.run(read_figure(fp, [_png(), _png()], "το σχήμα της λύσης", "Δίνεται τρίγωνο", "κ"))
    assert fig["crop_page"] == 1 and fig["spec_json"] == spec and fig["checked"]
    assert "το σχήμα της λύσης" in fp.calls[0]["text"] and fp.calls[0]["n"] == 2
    none = _FakeProvider([{"found": False, "crop_page": 0, "crop": [0, 0, 1, 1], "description": "", "why": "μόνο κείμενο"}])
    assert "Δεν βρέθηκε" in asyncio.run(read_figure(none, [_png()], "", "", "κ"))["error"]


# ------------------------------------------------------------------ 1.7.7: έλεγχος ενημερώσεων σε κάθε υπολογιστή

def test_update_contexts_prefer_windows_verification():
    from app import updates

    names = [n for n, _ in updates._contexts()]
    assert names[0] == "Windows" and "Python" in names and names[-1] == "certifi"


def test_update_errors_are_explained():
    import socket
    import ssl
    import urllib.error

    from app.updates import _explain

    def msg(reason):
        return str(_explain(urllib.error.URLError(reason), ["x"]))

    cert = ssl.SSLCertVerificationError(1, "[SSL: CERTIFICATE_VERIFY_FAILED] certificate verify failed: unable to get local issuer certificate")
    assert "πιστοποιητικό" in msg(cert) and "antivirus" in msg(cert)
    old = ssl.SSLCertVerificationError(1, "certificate verify failed: certificate has expired")
    assert "ώρα" in msg(old)
    assert "DNS" in msg(socket.gaierror(11001, "getaddrinfo failed"))
    assert "εγκαίρως" in msg(TimeoutError("timed out"))
    assert "firewall" in msg(PermissionError(10013, "An attempt was made to access a socket in a way forbidden by its access permissions"))
    assert "internet" in msg(OSError("άγνωστο"))
    e = _explain(urllib.error.URLError(cert), ["Windows: SSLError: x", "Python: SSLError: y"])
    assert "Windows" in e.detail and "Python" in e.detail


def test_update_open_falls_back_only_on_certificate_errors(monkeypatch):
    import io
    import socket
    import ssl
    import urllib.error
    import urllib.request

    from app import updates

    calls = []

    def fake_urlopen(req, timeout=10, context=None):
        calls.append(context)
        if len(calls) == 1:
            raise urllib.error.URLError(ssl.SSLCertVerificationError(1, "certificate verify failed"))
        return io.BytesIO(b'{"tag_name": "v9.9.9", "assets": []}')

    monkeypatch.setattr(urllib.request, "urlopen", fake_urlopen)
    assert updates.check("EDaskal/MATHEMATICS", "1.0.0")["latest"] == "9.9.9"
    assert len(calls) == 2  # ο 1ος τρόπος απέτυχε στο πιστοποιητικό → δοκιμάστηκε ο 2ος

    calls.clear()

    def no_net(req, timeout=10, context=None):
        calls.append(context)
        raise urllib.error.URLError(socket.gaierror(11001, "getaddrinfo failed"))

    monkeypatch.setattr(urllib.request, "urlopen", no_net)
    with pytest.raises(updates.UpdateError) as ei:
        updates.check("EDaskal/MATHEMATICS", "1.0.0")
    assert len(calls) == 1 and "DNS" in str(ei.value)  # χωρίς δίκτυο δεν έχει νόημα να δοκιμάσει άλλο τρόπο


def test_bridge_update_error_has_detail_and_page(monkeypatch, tmp_path):
    monkeypatch.setenv("APPDATA", str(tmp_path))
    from app import updates
    from app.bridge import Api

    def boom(url, timeout=10):
        raise updates.UpdateError("Δεν επαληθεύτηκε το πιστοποιητικό", "Windows: SSLError: x")

    monkeypatch.setattr(updates, "_get_json", boom)
    r = Api().check_updates()
    assert r["ok"] is False and r["detail"].startswith("Windows") and r["page"].endswith("/releases/latest")
